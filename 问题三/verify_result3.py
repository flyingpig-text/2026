# -*- coding: utf-8 -*-
"""
问题3结果最终质量检查。

检查内容：
1. result3.xlsx的计划购电量和调整购电量均为334天×144时段；
2. 充放电量、紧急购电量与逐10分钟明细逐项一致；
3. 计划购电费、调整费用和紧急购电费可由明细独立反算；
4. 表1、表2、表3的工作表结构和指定日期数值正确；
5. figures目录中的PNG文件均可正常解码且不是空白图。
"""

from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from PIL import Image


TOL = 1e-6
TARGET_DATES = (
    pd.Timestamp("2025-03-20").date(),
    pd.Timestamp("2025-06-21").date(),
    pd.Timestamp("2025-09-23").date(),
    pd.Timestamp("2025-12-21").date(),
)
FOUR_HOUR_BLOCKS = (
    "0:00-4:00",
    "4:00-8:00",
    "8:00-12:00",
    "12:00-16:00",
    "16:00-20:00",
    "20:00-24:00",
)


class Verification:
    """收集全部检查结果，最后统一输出。"""

    def __init__(self) -> None:
        self.passed: list[str] = []
        self.failed: list[str] = []

    def check(self, condition: bool, message: str, detail: str = "") -> None:
        """记录一项通过或失败，不让单个失败阻断其余检查。"""
        if condition:
            self.passed.append(message)
            print(f"[通过] {message}")
        else:
            text = message if not detail else f"{message}；{detail}"
            self.failed.append(text)
            print(f"[失败] {text}")

    def close(self) -> None:
        """输出汇总并在存在失败项时返回非零退出码。"""
        print("\n质量检查汇总：")
        print(f"通过：{len(self.passed)} 项")
        print(f"失败：{len(self.failed)} 项")
        if self.failed:
            print("\n失败明细：")
            for item in self.failed:
                print(f"- {item}")
            raise SystemExit(1)


def close_enough(actual: float, expected: float, tolerance: float = TOL) -> bool:
    """执行考虑浮点误差的数值比较。"""
    return abs(float(actual) - float(expected)) <= tolerance


def arrays_close(actual: np.ndarray, expected: np.ndarray) -> bool:
    """执行数组整体比较。"""
    return (
        actual.shape == expected.shape
        and bool(np.all(np.isfinite(actual)))
        and float(np.max(np.abs(actual - expected))) <= TOL
    )


def read_wide_sheet(sheet) -> tuple[list[str], list, np.ndarray, np.ndarray, np.ndarray]:
    """读取result3的日期×144时段宽表、日购电量和日费用。"""
    rows = list(sheet.iter_rows(values_only=True))
    headers = [str(value) for value in rows[0]]
    data_rows = rows[1:]
    dates = [pd.Timestamp(row[0]).date() for row in data_rows]
    values = np.asarray(
        [[float(value) for value in row[1:145]] for row in data_rows],
        dtype=float,
    )
    daily_quantity = np.asarray([float(row[145]) for row in data_rows], dtype=float)
    daily_cost = np.asarray([float(row[146]) for row in data_rows], dtype=float)
    return headers, dates, values, daily_quantity, daily_cost


def main() -> None:
    """执行全部结果质量检查。"""
    script_dir = Path(__file__).resolve().parent
    result_dir = script_dir / "results"
    detail_path = result_dir / "tables" / "逐10分钟计划调整明细.csv"
    result3_path = result_dir / "result3.xlsx"
    table1_path = result_dir / "表1_指定日期购电量.xlsx"
    table2_path = result_dir / "表2_指定日期充放电量.xlsx"
    table3_path = result_dir / "表3_指定日期紧急购电量.xlsx"
    figures_dir = result_dir / "figures"

    verification = Verification()
    required_files = (
        detail_path,
        result3_path,
        table1_path,
        table2_path,
        table3_path,
    )
    for path in required_files:
        verification.check(path.is_file(), f"文件存在：{path.name}")

    detail = pd.read_csv(detail_path, encoding="utf-8-sig", parse_dates=["日期"])
    detail["日期"] = pd.to_datetime(detail["日期"])
    detail = detail.sort_values(["日期", "时段序号"]).reset_index(drop=True)
    detail_dates = sorted(detail["日期"].dt.date.unique())
    expected_dates = [
        value.date()
        for value in pd.date_range("2025-02-01", "2025-12-31", freq="D")
    ]
    verification.check(len(detail_dates) == 334, "明细覆盖334天")
    verification.check(detail_dates == expected_dates, "明细日期为2025-02-01至2025-12-31")
    periods_per_day = detail.groupby("日期").size().unique().tolist()
    verification.check(periods_per_day == [144], "明细每天恰有144个10分钟时段")
    verification.check(len(detail) == 334 * 144, "明细总行数为334×144")
    verification.check(int(detail.isna().sum().sum()) == 0, "逐10分钟明细不存在空值")

    detail_by_day = {
        current_date: day.sort_values("时段序号")
        for current_date, day in detail.groupby(detail["日期"].dt.date)
    }

    workbook = load_workbook(result3_path, read_only=True, data_only=True)
    expected_sheets = ["计划购电量", "调整购电量", "充放电量", "紧急购电量"]
    verification.check(workbook.sheetnames == expected_sheets, "result3工作表名称和顺序正确")

    plan_headers, plan_dates, plan_matrix, plan_quantity, plan_cost = read_wide_sheet(
        workbook["计划购电量"]
    )
    adjusted_headers, adjusted_dates, adjusted_matrix, adjusted_quantity, adjusted_cost = (
        read_wide_sheet(workbook["调整购电量"])
    )
    verification.check(len(plan_headers) == 147, "计划购电量工作表有147列")
    verification.check(len(adjusted_headers) == 147, "调整购电量工作表有147列")
    verification.check(plan_dates == expected_dates, "计划购电量工作表日期完整")
    verification.check(adjusted_dates == expected_dates, "调整购电量工作表日期完整")
    verification.check(plan_matrix.shape == (334, 144), "计划购电量矩阵为334天×144时段")
    verification.check(adjusted_matrix.shape == (334, 144), "调整购电量矩阵为334天×144时段")
    verification.check(bool(np.all(np.isfinite(plan_matrix))), "计划购电量矩阵无空值")
    verification.check(bool(np.all(np.isfinite(adjusted_matrix))), "调整购电量矩阵无空值")

    expected_plan = np.vstack(
        [
            detail_by_day[current_date]["计划购电量_kWh"].to_numpy(dtype=float)
            for current_date in expected_dates
        ]
    )
    expected_adjusted = np.vstack(
        [
            detail_by_day[current_date]["调整购电量_kWh"].to_numpy(dtype=float)
            for current_date in expected_dates
        ]
    )
    verification.check(
        arrays_close(plan_matrix, expected_plan),
        "计划购电量工作表逐时段与明细一致",
    )
    verification.check(
        arrays_close(adjusted_matrix, expected_adjusted),
        "调整购电量工作表逐时段与明细一致",
    )
    verification.check(
        arrays_close(
            plan_quantity,
            np.asarray([expected_plan[index].sum() for index in range(334)]),
        ),
        "计划购电量日合计与明细一致",
    )
    verification.check(
        arrays_close(
            adjusted_quantity,
            np.asarray([expected_adjusted[index].sum() for index in range(334)]),
        ),
        "调整购电量日合计与明细一致",
    )
    expected_plan_cost = np.asarray(
        [
            detail_by_day[current_date]["计划购电费_元"].sum()
            for current_date in expected_dates
        ],
        dtype=float,
    )
    expected_adjustment_cost = np.asarray(
        [
            detail_by_day[current_date]["调整费用_元"].sum()
            for current_date in expected_dates
        ],
        dtype=float,
    )
    verification.check(
        arrays_close(plan_cost, expected_plan_cost),
        "计划购电费与明细逐日反算一致",
    )
    verification.check(
        arrays_close(adjusted_cost, expected_adjustment_cost),
        "调整费用与明细逐日反算一致",
    )

    charge_rows = list(workbook["充放电量"].iter_rows(values_only=True))
    emergency_rows = list(workbook["紧急购电量"].iter_rows(values_only=True))
    verification.check(
        len(charge_rows) == 1 + 334 * 6,
        "充放电量工作表为表头加334天×6个4小时区间",
        f"实际行数={len(charge_rows)}",
    )
    verification.check(
        charge_rows[0] == ("日期", "时间段", "充电量(kWh)", "放电量(kWh)", "时刻", "储电量(kWh)"),
        "充放电量表头格式正确",
    )
    charge_ok = True
    soc_ok = True
    charge_detail = ""
    for day_index, current_date in enumerate(expected_dates):
        block = charge_rows[1 + day_index * 6 : 1 + (day_index + 1) * 6]
        day = detail_by_day[current_date]
        expected_charge = np.asarray(
            [
                day["充电量_kWh"].iloc[index : index + 24].sum()
                for index in range(0, 144, 24)
            ],
            dtype=float,
        )
        expected_discharge = np.asarray(
            [
                day["放电量_kWh"].iloc[index : index + 24].sum()
                for index in range(0, 144, 24)
            ],
            dtype=float,
        )
        actual_charge = np.asarray([float(row[2]) for row in block], dtype=float)
        actual_discharge = np.asarray([float(row[3]) for row in block], dtype=float)
        block_names = [row[1] for row in block]
        if (
            block_names != list(FOUR_HOUR_BLOCKS)
            or not arrays_close(actual_charge, expected_charge)
            or not arrays_close(actual_discharge, expected_discharge)
        ):
            charge_ok = False
            charge_detail = f"首个不一致日期={current_date}"
            break
        expected_initial_soc = 6000.0
        expected_final_soc = float(day.iloc[-1]["时段末储电量_kWh"])
        if (
            not close_enough(float(block[0][5]), expected_initial_soc)
            or not close_enough(float(block[1][5]), expected_final_soc)
            or block[0][4] != "0:00"
            or block[1][4] != "24:00"
        ):
            soc_ok = False
            break
    verification.check(charge_ok, "充放电量工作表与明细4小时聚合一致", charge_detail)
    verification.check(soc_ok, "充放电量工作表0:00和24:00储电量正确")

    emergency_detail = detail[detail["紧急购电量_kWh"] > 1e-8].sort_values(
        ["日期", "时段序号"]
    )
    verification.check(
        len(emergency_rows) == 1 + len(emergency_detail),
        "紧急购电量事件行数完整",
        f"工作表={len(emergency_rows) - 1}，明细={len(emergency_detail)}",
    )
    emergency_ok = True
    emergency_message = ""
    for row, expected in zip(emergency_rows[1:], emergency_detail.itertuples(index=False)):
        if (
            pd.Timestamp(row[0]).date() != expected.日期.date()
            or row[1] != expected.时段
            or not close_enough(float(row[2]), float(expected.紧急购电量_kWh))
        ):
            emergency_ok = False
            emergency_message = str(expected.日期.date())
            break
    verification.check(emergency_ok, "紧急购电日期、时段和电量与明细一致", emergency_message)

    total_plan_quantity = float(detail["计划购电量_kWh"].sum())
    total_adjusted_quantity = float(detail["调整购电量_kWh"].sum())
    total_emergency_quantity = float(detail["紧急购电量_kWh"].sum())
    total_plan_cost = float(detail["计划购电费_元"].sum())
    total_adjustment_cost = float(detail["调整费用_元"].sum())
    total_emergency_cost = float(detail["紧急购电费_元"].sum())
    total_cost = total_plan_cost + total_adjustment_cost + total_emergency_cost
    verification.check(
        close_enough(plan_matrix.sum(), total_plan_quantity),
        "全年计划购电量由逐时段明细反算一致",
    )
    verification.check(
        close_enough(adjusted_matrix.sum(), total_adjusted_quantity),
        "全年调整购电量由逐时段明细反算一致",
    )
    verification.check(
        close_enough(
            sum(float(row[2]) for row in emergency_rows[1:]),
            total_emergency_quantity,
        ),
        "全年紧急购电量由事件明细反算一致",
    )
    verification.check(
        close_enough(plan_cost.sum(), total_plan_cost),
        "全年计划购电费由逐日明细反算一致",
    )
    verification.check(
        close_enough(adjusted_cost.sum(), total_adjustment_cost),
        "全年调整费用由逐日明细反算一致",
    )
    verification.check(
        close_enough(total_cost, 16609954.260720413, tolerance=1e-5),
        "三类费用之和等于总费用",
        f"反算总费用={total_cost:.6f} 元",
    )

    table1_rows = list(
        load_workbook(table1_path, read_only=True, data_only=True)["表1_指定日期购电量"].iter_rows(
            values_only=True
        )
    )
    table1_ok = len(table1_rows) == 5 and len(table1_rows[0]) == 15
    table1_message = f"行列数={len(table1_rows)}×{len(table1_rows[0])}"
    if table1_ok:
        intervals = (
            "10:00-10:10",
            "12:00-12:10",
            "14:00-14:10",
            "16:00-16:10",
            "18:00-18:10",
            "20:00-20:10",
        )
        for row, target in zip(table1_rows[1:], TARGET_DATES):
            day = detail_by_day[target]
            expected_values = [
                float(day[day["时段"] == interval].iloc[0]["计划购电量_kWh"])
                for interval in intervals
            ]
            actual_values = [float(row[2 + 2 * index]) for index in range(6)]
            if (
                pd.Timestamp(row[0]).date() != target
                or not arrays_close(np.asarray(actual_values), np.asarray(expected_values))
                or not close_enough(float(row[13]), float(day["计划购电量_kWh"].sum()))
                or not close_enough(float(row[14]), float(day["计划购电费_元"].sum()))
            ):
                table1_ok = False
                table1_message = f"不一致日期={target}"
                break
    verification.check(table1_ok, "表1格式及指定日期数值正确", table1_message)

    table2_rows = list(
        load_workbook(table2_path, read_only=True, data_only=True)["表2_指定日期充放电量"].iter_rows(
            values_only=True
        )
    )
    table2_ok = len(table2_rows) == 25 and len(table2_rows[0]) == 6
    table2_message = f"行列数={len(table2_rows)}×{len(table2_rows[0])}"
    if table2_ok:
        for day_index, target in enumerate(TARGET_DATES):
            block = table2_rows[1 + day_index * 6 : 1 + (day_index + 1) * 6]
            day = detail_by_day[target]
            expected_charge = np.asarray(
                [
                    day["充电量_kWh"].iloc[index : index + 24].sum()
                    for index in range(0, 144, 24)
                ],
                dtype=float,
            )
            expected_discharge = np.asarray(
                [
                    day["放电量_kWh"].iloc[index : index + 24].sum()
                    for index in range(0, 144, 24)
                ],
                dtype=float,
            )
            if (
                pd.Timestamp(block[0][0]).date() != target
                or [row[1] for row in block] != list(FOUR_HOUR_BLOCKS)
                or not arrays_close(
                    np.asarray([float(row[2]) for row in block]),
                    expected_charge,
                )
                or not arrays_close(
                    np.asarray([float(row[3]) for row in block]),
                    expected_discharge,
                )
                or not close_enough(float(block[0][5]), 6000.0)
                or not close_enough(
                    float(block[1][5]),
                    float(day.iloc[-1]["时段末储电量_kWh"]),
                )
            ):
                table2_ok = False
                table2_message = f"不一致日期={target}"
                break
    verification.check(table2_ok, "表2格式及指定日期充放电和储电量正确", table2_message)

    table3_sheet = load_workbook(table3_path, read_only=True, data_only=True)[
        "表3_紧急购电量"
    ]
    table3_rows = list(table3_sheet.iter_rows(values_only=True))
    table3_ok = (
        len(table3_rows) >= 3
        and len(table3_rows[1]) == 9
        and all(
            [table3_rows[0][1 + 2 * index], table3_rows[1][1 + 2 * index]]
            == [target.strftime("%Y.%m.%d"), "时间段"]
            for index, target in enumerate(TARGET_DATES)
        )
    )
    table3_message = f"行列数={table3_sheet.max_row}×{table3_sheet.max_column}"
    if table3_ok:
        for index, target in enumerate(TARGET_DATES):
            expected = emergency_detail[
                emergency_detail["日期"].dt.date == target
            ].reset_index(drop=True)
            actual = [
                row
                for row in table3_rows[2:]
                if row[1 + 2 * index] is not None
            ]
            if len(actual) != len(expected):
                table3_ok = False
                table3_message = f"事件数量不一致日期={target}"
                break
            for actual_row, expected_row in zip(actual, expected.itertuples(index=False)):
                if (
                    actual_row[1 + 2 * index] != expected_row.时段
                    or not close_enough(
                        float(actual_row[2 + 2 * index]),
                        float(expected_row.紧急购电量_kWh),
                    )
                ):
                    table3_ok = False
                    table3_message = f"数值不一致日期={target}"
                    break
            if not table3_ok:
                break
    verification.check(table3_ok, "表3四日期并排格式及紧急购电数值正确", table3_message)

    image_files = sorted(figures_dir.glob("*.png"))
    verification.check(len(image_files) == 3, "图像目录包含3张PNG结果图")
    for path in image_files:
        try:
            with Image.open(path) as image:
                image.verify()
            with Image.open(path).convert("RGB") as image:
                array = np.asarray(image, dtype=np.uint8)
                nonwhite_fraction = float(np.mean(np.any(array < 245, axis=2)))
                unique_colors = len(np.unique(array.reshape(-1, 3), axis=0))
            verification.check(
                image.width > 0
                and image.height > 0
                and nonwhite_fraction > 0.01
                and unique_colors > 10,
                f"图像可解码且非空白：{path.name}",
                f"尺寸={image.width}×{image.height}，非白像素={nonwhite_fraction:.3%}",
            )
        except Exception as error:  # pragma: no cover - 仅用于最终质量检查报告
            verification.check(False, f"图像可正常读取：{path.name}", str(error))

    print("\n关键反算结果：")
    print(f"计划购电量 = {total_plan_quantity:.6f} kWh")
    print(f"调整购电量 = {total_adjusted_quantity:.6f} kWh")
    print(f"紧急购电量 = {total_emergency_quantity:.6f} kWh")
    print(f"计划购电费 = {total_plan_cost:.6f} 元")
    print(f"调整费用 = {total_adjustment_cost:.6f} 元")
    print(f"紧急购电费 = {total_emergency_cost:.6f} 元")
    print(f"总费用 = {total_cost:.6f} 元")
    verification.close()


if __name__ == "__main__":
    main()

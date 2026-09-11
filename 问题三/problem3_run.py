# -*- coding: utf-8 -*-
"""2026 C题问题3：滚动光伏预报、计划购电与调整购电求解程序。"""

from __future__ import annotations

import argparse
import json
import os
import sys
from datetime import date
from pathlib import Path

import matplotlib

MPL_CACHE_DIR = Path(__file__).resolve().parent / ".cache" / "matplotlib"
MPL_CACHE_DIR.mkdir(parents=True, exist_ok=True)
os.environ.setdefault("MPLCONFIGDIR", str(MPL_CACHE_DIR))
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from problem3_core import (
    DT_H,
    OUTPUT_END,
    OUTPUT_START,
    SETTLEMENT_MODES,
    T,
    TARGET_DATES,
    aggregate_forecast_scenarios,
    dataframe_row_for_day,
    load_problem2_module,
    locate_inputs,
    merge_contiguous_emergency_events,
    prepare_actual_data,
    print_quantity_checks,
    read_attachment1_load_energy,
    read_attachment3,
    summarize_specified_dates,
    validate_result_detail,
    write_official_result,
    write_specified_date_workbook,
)
from problem3_algorithm import run_rolling_day
from problem3_scenarios import (
    build_causal_load_forecast,
    build_day_scenario_windows,
)


plt.rcParams["font.sans-serif"] = [
    "Microsoft YaHei",
    "SimHei",
    "Noto Sans CJK SC",
    "Arial Unicode MS",
    "DejaVu Sans",
]
plt.rcParams["axes.unicode_minus"] = False


def configure_console() -> None:
    """统一控制台编码。"""
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    if hasattr(sys.stderr, "reconfigure"):
        sys.stderr.reconfigure(encoding="utf-8")


def parse_args() -> argparse.Namespace:
    """解析费用口径和输出目录参数。"""
    parser = argparse.ArgumentParser(description="2026 C题问题3滚动优化")
    parser.add_argument(
        "--settlement-mode",
        choices=SETTLEMENT_MODES,
        default="plan_full",
        help=(
            "plan_full：计划购电量始终按正常电价结算；"
            "actual_base：正常电价只结算min(计划购电量,调整购电量)。"
        ),
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=None,
        help="结果目录；默认写入问题三/results。",
    )
    parser.add_argument(
        "--scenarios",
        type=int,
        default=5,
        help="每个决策时刻使用的历史误差情景数，默认5。",
    )
    parser.add_argument(
        "--scenario-lookback-days",
        type=int,
        default=30,
        help="历史误差情景回看天数，默认30。",
    )
    parser.add_argument(
        "--window-days",
        type=int,
        default=3,
        help="情景滚动优化窗口天数，默认3。",
    )
    parser.add_argument(
        "--terminal-value-factor",
        type=float,
        default=1.0,
        help="窗口末端单位储能价值倍率，默认1.0。",
    )
    parser.add_argument(
        "--scenario-time-limit",
        type=float,
        default=60.0,
        help="单个情景窗口MILP时间上限，单位秒，默认60。",
    )
    return parser.parse_args()


def output_dates(data: pd.DataFrame) -> list[date]:
    """返回问题3要求输出的日期。"""
    return sorted(
        current_date
        for current_date in data["日期"].dt.date.unique()
        if OUTPUT_START <= current_date <= OUTPUT_END
    )


def run_baseline_benchmark(
    p2,
    data: pd.DataFrame,
    storage,
) -> pd.DataFrame:
    """
    运行简单基准算例：储能不动作，实际净负荷全部由计划购电满足。

    该基准不涉及滚动调整，只用于确认附件读取、电量换算和电能平衡。
    """
    rows: list[dict[str, object]] = []
    print("步骤1：运行储能不动作基准算例。")
    for target in TARGET_DATES:
        day = data[data["日期"].dt.date == target].sort_values("时段序号")
        load = day["小区负载电量_kWh"].to_numpy(dtype=float)
        pv = day["光伏实际电量_kWh"].to_numpy(dtype=float)
        price = day["电价_元每kWh"].to_numpy(dtype=float)
        dispatch = p2.baseline_day(load, pv, price, storage.initial_kwh)
        balance_error = (
            dispatch.planned_purchase_kwh
            + pv
            + dispatch.discharge_kwh
            - load
            - dispatch.charge_kwh
            - dispatch.curtail_kwh
        )
        rows.append(
            {
                "日期": target,
                "计划购电量_kWh": float(dispatch.planned_purchase_kwh.sum()),
                "购电费_元": float(dispatch.total_cost_yuan),
                "最大电能平衡残差_kWh": float(np.max(np.abs(balance_error))),
            }
        )
        print(
            f"{target}：计划购电量="
            f"{dispatch.planned_purchase_kwh.sum():.6f} kWh，"
            f"购电费={dispatch.total_cost_yuan:.6f} 元，"
            f"最大平衡残差={np.max(np.abs(balance_error)):.3e} kWh。"
        )
    result = pd.DataFrame(rows)
    result["日期"] = pd.to_datetime(result["日期"])
    return result


def solve_problem3(
    p2,
    data: pd.DataFrame,
    forecasts: dict[date, dict[int, np.ndarray]],
    storage,
    settlement_mode: str = "plan_full",
    decision_price_by_date: dict[date, np.ndarray] | None = None,
    fallback_load_profile_kwh: np.ndarray | None = None,
    scenario_count: int = 5,
    scenario_lookback_days: int = 30,
    window_days: int = 3,
    terminal_soc_value_yuan_per_kwh: float = 0.0,
    scenario_time_limit_s: float = 60.0,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """从1月1日开始连续运行，输出2月1日至12月31日的滚动结果。"""
    detail_rows: list[dict[str, object]] = []
    daily_rows: list[dict[str, object]] = []
    scenario_rows: list[dict[str, object]] = []
    all_dates = sorted(data["日期"].dt.date.unique())
    output_period_dates = output_dates(data)
    current_soc_kwh = float(storage.initial_kwh)
    print(
        f"问题3开始连续求解：{len(all_dates)}天，"
        f"输出{len(output_period_dates)}天，每天144个10分钟时段。"
    )
    for number, current_date in enumerate(all_dates, start=1):
        day = data[data["日期"].dt.date == current_date].sort_values("时段序号")
        decision_load = build_causal_load_forecast(
            data,
            current_date,
            fallback_load_profile_kwh,
        )
        decision_price = (
            decision_price_by_date[current_date]
            if decision_price_by_date is not None
            else day["电价_元每kWh"].to_numpy(dtype=float)
        )
        scenario_windows = build_day_scenario_windows(
            data,
            forecasts,
            current_date,
            decision_price,
            fallback_load_profile_kwh,
            scenario_count=scenario_count,
            lookback_days=scenario_lookback_days,
            window_days=window_days,
        )
        rolling = run_rolling_day(
            load_energy_kwh=day["小区负载电量_kWh"].to_numpy(dtype=float),
            actual_pv_energy_kwh=day["光伏实际电量_kWh"].to_numpy(dtype=float),
            price_yuan_per_kwh=day["电价_元每kWh"].to_numpy(dtype=float),
            forecast_by_hour=forecasts[current_date],
            storage=storage,
            initial_soc_kwh=current_soc_kwh,
            settlement_mode=settlement_mode,
            decision_price_yuan_per_kwh=decision_price,
            forecast_load_energy_kwh=decision_load,
            scenario_windows_by_hour=scenario_windows,
            live_storage_execution=True,
            terminal_soc_value_yuan_per_kwh=terminal_soc_value_yuan_per_kwh,
            scenario_time_limit_s=scenario_time_limit_s,
        )
        result = rolling.as_dict() if hasattr(rolling, "as_dict") else rolling
        current_soc_kwh = float(result["soc_kwh"][-1])
        if not (OUTPUT_START <= current_date <= OUTPUT_END):
            continue
        rows, daily = dataframe_row_for_day(current_date, data, result)
        detail_rows.extend(rows)
        daily_rows.append(daily)
        for scenario in result["scenarios"]:
            scenario_rows.append({"日期": current_date, **scenario})
        if len(daily_rows) % 30 == 0 or current_date == OUTPUT_END:
            print(
                f"问题3输出完成 {len(daily_rows):>3}/{len(output_period_dates)} 天：{current_date}，"
                f"调整购电={daily['调整购电量_kWh']:.6f} kWh，"
                f"紧急购电={daily['紧急购电量_kWh']:.6f} kWh，"
                f"总费用={daily['总费用_元']:.6f} 元。"
            )

    detail = pd.DataFrame(detail_rows)
    daily = pd.DataFrame(daily_rows)
    scenarios = pd.DataFrame(scenario_rows)
    detail["日期"] = pd.to_datetime(detail["日期"])
    daily["日期"] = pd.to_datetime(daily["日期"])
    scenarios["日期"] = pd.to_datetime(scenarios["日期"])
    return detail, daily, scenarios


def write_table1_excel(
    detail: pd.DataFrame,
    output_path: Path,
) -> None:
    """按表1格式写出计划购电量、最终购电量和调整对照。"""
    intervals = (
        "10:00-10:10",
        "12:00-12:10",
        "14:00-14:10",
        "16:00-16:10",
        "18:00-18:10",
        "20:00-20:10",
    )
    workbook = Workbook()
    plan_sheet = workbook.active
    plan_sheet.title = "表1_计划购电量"
    headers = ["日期"]
    for interval in intervals:
        headers.extend([f"{interval}时间段", f"{interval}购电量(kWh)"])
    headers.extend(["全天购电量(kWh)", "全天购电费(元)"])
    plan_sheet.append(headers)

    for target in TARGET_DATES:
        day = detail[detail["日期"].dt.date == target].sort_values("时段序号")
        record: list[object] = [target.strftime("%Y-%m-%d")]
        for interval in intervals:
            selected = day[day["时段"] == interval]
            if len(selected) != 1:
                raise ValueError(f"{target}的{interval}记录不唯一。")
            record.extend(
                [
                    interval,
                    float(selected.iloc[0]["计划购电量_kWh"]),
                ]
            )
        record.extend(
            [
                float(day["计划购电量_kWh"].sum()),
                float(day["计划购电费_元"].sum()),
            ]
        )
        plan_sheet.append(record)

    final_sheet = workbook.create_sheet("表1_最终购电量")
    final_sheet.append(headers)
    for target in TARGET_DATES:
        day = detail[detail["日期"].dt.date == target].sort_values("时段序号")
        record = [target.strftime("%Y-%m-%d")]
        for interval in intervals:
            selected = day[day["时段"] == interval]
            if len(selected) != 1:
                raise ValueError(f"{target}的{interval}记录不唯一。")
            record.extend(
                [
                    interval,
                    float(selected.iloc[0]["调整购电量_kWh"]),
                ]
            )
        total_cost = float(
            day[
                ["计划购电费_元", "调整费用_元", "紧急购电费_元"]
            ].to_numpy(dtype=float).sum()
        )
        record.extend(
            [
                float(day["调整购电量_kWh"].sum()),
                total_cost,
            ]
        )
        final_sheet.append(record)

    comparison_sheet = workbook.create_sheet("计划调整对照")
    comparison_sheet.append(
        [
            "日期",
            "时间段",
            "电价(元/kWh)",
            "计划购电量(kWh)",
            "最终购电量(kWh)",
            "调整净变化(kWh)",
        ]
    )
    for target in TARGET_DATES:
        day = detail[detail["日期"].dt.date == target].sort_values("时段序号")
        for interval in intervals:
            selected = day[day["时段"] == interval].iloc[0]
            plan_value = float(selected["计划购电量_kWh"])
            adjusted_value = float(selected["调整购电量_kWh"])
            comparison_sheet.append(
                [
                    target.strftime("%Y-%m-%d"),
                    interval,
                    float(selected["电价_元每kWh"]),
                    plan_value,
                    adjusted_value,
                    adjusted_value - plan_value,
                ]
            )

    for worksheet in workbook.worksheets:
        for cell in worksheet[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center", vertical="center")
        for column in range(1, worksheet.max_column + 1):
            worksheet.column_dimensions[get_column_letter(column)].width = 22
        worksheet.freeze_panes = "B2"
    for cell in plan_sheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    workbook.save(output_path)
    workbook.close()


def write_table2_excel(
    p2,
    detail: pd.DataFrame,
    storage,
    output_path: Path,
) -> None:
    """按表2格式写出指定日期的4小时充放电量和首末储电量。"""
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "表2_指定日期充放电量"
    worksheet.append(["日期", "时间段", "充电量(kWh)", "放电量(kWh)", "时刻", "储电量(kWh)"])

    for target in TARGET_DATES:
        day = detail[detail["日期"].dt.date == target].sort_values("时段序号")
        charge_blocks = p2.aggregate_four_hour(
            day["充电量_kWh"].to_numpy(dtype=float)
        )
        discharge_blocks = p2.aggregate_four_hour(
            day["放电量_kWh"].to_numpy(dtype=float)
        )
        start_row = worksheet.max_row + 1
        for index, block in enumerate(p2.FOUR_HOUR_BLOCKS):
            worksheet.append(
                [
                    target.strftime("%Y-%m-%d") if index == 0 else None,
                    block,
                    charge_blocks[index],
                    discharge_blocks[index],
                    None,
                    None,
                ]
            )
        worksheet.cell(start_row, 5, "0:00")
        worksheet.cell(
            start_row,
            6,
            float(day.iloc[0]["时段初储电量_kWh"]),
        )
        worksheet.cell(start_row + 1, 5, "24:00")
        worksheet.cell(
            start_row + 1,
            6,
            float(day.iloc[-1]["时段末储电量_kWh"]),
        )

    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for column in range(1, 7):
        worksheet.column_dimensions[get_column_letter(column)].width = 20
    worksheet.freeze_panes = "A2"
    workbook.save(output_path)
    workbook.close()


def write_paper_tables(
    p2,
    detail: pd.DataFrame,
    storage,
    output_dir: Path,
) -> dict[str, object]:
    """写出论文要求的表1、表2、表3格式工作簿。"""
    table1 = output_dir / "表1_指定日期购电量.xlsx"
    table2 = output_dir / "表2_指定日期充放电量.xlsx"
    table3 = output_dir / "表3_指定日期紧急购电量.xlsx"
    write_table1_excel(detail, table1)
    write_table2_excel(p2, detail, storage, table2)
    table3_data = merge_contiguous_emergency_events(detail)
    p2.write_table3_excel(table3_data, table3)
    return {
        "表1": table1,
        "表2": table2,
        "表3": table3,
        "表3明细": table3_data,
    }


def run_forecast_sensitivity(
    p2,
    data: pd.DataFrame,
    forecasts: dict[date, dict[int, np.ndarray]],
    storage,
    settlement_mode: str = "plan_full",
    initial_soc_by_date: dict[date, float] | None = None,
    scenario_count: int = 5,
    scenario_lookback_days: int = 30,
    window_days: int = 3,
    terminal_soc_value_yuan_per_kwh: float = 0.0,
    scenario_time_limit_s: float = 60.0,
    fallback_profile_kwh: np.ndarray | None = None,
) -> pd.DataFrame:
    """对指定日期进行预报整体缩放灵敏度分析。"""
    rows: list[dict[str, object]] = []
    for target in TARGET_DATES:
        day = data[data["日期"].dt.date == target].sort_values("时段序号")
        initial_soc = (
            storage.initial_kwh
            if initial_soc_by_date is None
            else initial_soc_by_date[target]
        )
        decision_load = build_causal_load_forecast(
            data,
            target,
            fallback_profile_kwh
            if fallback_profile_kwh is not None
            else day["小区负载电量_kWh"].to_numpy(dtype=float),
        )
        for scale in (0.90, 0.95, 1.00, 1.05, 1.10):
            scenario_windows = build_day_scenario_windows(
                data,
                forecasts,
                target,
                day["电价_元每kWh"].to_numpy(dtype=float),
                fallback_profile_kwh=(
                    fallback_profile_kwh
                    if fallback_profile_kwh is not None
                    else day["小区负载电量_kWh"].to_numpy(dtype=float)
                ),
                scenario_count=scenario_count,
                lookback_days=scenario_lookback_days,
                window_days=window_days,
                forecast_scale=scale,
            )
            rolling = run_rolling_day(
                load_energy_kwh=day["小区负载电量_kWh"].to_numpy(dtype=float),
                actual_pv_energy_kwh=day["光伏实际电量_kWh"].to_numpy(dtype=float),
                price_yuan_per_kwh=day["电价_元每kWh"].to_numpy(dtype=float),
                forecast_by_hour=forecasts[target],
                storage=storage,
                initial_soc_kwh=initial_soc,
                forecast_scale=scale,
                settlement_mode=settlement_mode,
                forecast_load_energy_kwh=decision_load,
                scenario_windows_by_hour=scenario_windows,
                live_storage_execution=True,
                terminal_soc_value_yuan_per_kwh=terminal_soc_value_yuan_per_kwh,
                scenario_time_limit_s=scenario_time_limit_s,
            )
            result = rolling.as_dict() if hasattr(rolling, "as_dict") else rolling
            rows.append(
                {
                    "日期": target,
                    "预报整体缩放比例": scale,
                    "计划购电量_kWh": float(result["plan_purchase_kwh"].sum()),
                    "调整购电量_kWh": float(
                        result["adjusted_purchase_kwh"].sum()
                    ),
                    "紧急购电量_kWh": float(
                        result["emergency_purchase_kwh"].sum()
                    ),
                    "总费用_元": float(result["total_cost_yuan"]),
                }
            )
    result = pd.DataFrame(rows)
    result["日期"] = pd.to_datetime(result["日期"])
    return result


def save_figure(fig, output_path: Path) -> None:
    """
    规范化输出路径并保存图片。

    Windows下若命令行传入重复反斜杠，Pillow可能抛出Errno 22。
    这里先用系统规范路径写入；若仍失败，再用正斜杠路径重试。
    """
    normalized = Path(os.path.normpath(os.fspath(output_path)))
    normalized.parent.mkdir(parents=True, exist_ok=True)
    try:
        fig.savefig(normalized, dpi=180)
    except OSError as first_error:
        if getattr(first_error, "errno", None) != 22:
            raise
        fig.savefig(normalized.as_posix(), dpi=180)


def plot_forecast_and_dispatch(detail: pd.DataFrame, output_path: Path) -> None:
    """绘制指定日期的实际光伏、预报和计划调整结果。"""
    hours = np.arange(1, T + 1) * DT_H
    fig, axes = plt.subplots(2, 2, figsize=(15, 9), sharex=True)
    for ax, target in zip(axes.flat, TARGET_DATES):
        day = detail[detail["日期"].dt.date == target].sort_values("时段序号")
        ax.plot(hours, day["光伏实际_kW"], label="光伏实际", color="#2ca02c")
        ax.plot(
            hours,
            day["光伏0时预报_kW"],
            label="0:00预报",
            color="#ff7f0e",
            linestyle="--",
        )
        ax.plot(
            hours,
            day["最终采用预报_kW"],
            label="最终采用预报",
            color="#9467bd",
            linestyle=":",
        )
        ax.step(
            hours,
            day["计划购电量_kWh"] / DT_H,
            where="post",
            label="计划购电等效功率",
            color="#1f77b4",
        )
        ax.step(
            hours,
            day["调整购电量_kWh"] / DT_H,
            where="post",
            label="调整购电等效功率",
            color="#d62728",
        )
        emergency_power = day["紧急购电量_kWh"].to_numpy(dtype=float) / DT_H
        if np.max(emergency_power) > 1e-8:
            ax.step(
                hours,
                emergency_power,
                where="post",
                label="紧急购电等效功率",
                color="#8c564b",
            )
        ax.set_title(target.strftime("%Y-%m-%d"))
        ax.set_xlim(0, 24)
        ax.set_xticks(np.arange(0, 25, 3))
        ax.set_ylabel("功率 (kW)")
        ax.grid(alpha=0.25)
        ax.legend(loc="upper left", fontsize=7.5)
    fig.suptitle("问题3指定日期：预报更新、计划购电与调整购电", fontsize=15)
    fig.tight_layout(rect=(0, 0, 1, 0.97))
    save_figure(fig, output_path)
    plt.close(fig)


def plot_storage(detail: pd.DataFrame, storage, output_path: Path) -> None:
    """绘制指定日期充放电功率和储电量。"""
    hours = np.arange(1, T + 1) * DT_H
    fig, axes = plt.subplots(2, 2, figsize=(15, 9), sharex=True)
    for ax, target in zip(axes.flat, TARGET_DATES):
        day = detail[detail["日期"].dt.date == target].sort_values("时段序号")
        charge_power = day["充电量_kWh"].to_numpy(dtype=float) / DT_H
        discharge_power = day["放电量_kWh"].to_numpy(dtype=float) / DT_H
        ax.step(hours, charge_power, where="post", label="充电功率", color="#2ca02c")
        ax.step(
            hours,
            -discharge_power,
            where="post",
            label="放电功率（负值）",
            color="#d62728",
        )
        ax.set_ylabel("充放电功率 (kW)")
        ax.set_ylim(-5100, 5100)
        ax.grid(alpha=0.25)
        ax2 = ax.twinx()
        soc = np.concatenate(
            (
                [storage.initial_kwh],
                day["时段末储电量_kWh"].to_numpy(dtype=float),
            )
        )
        ax2.plot(
            np.concatenate(([0.0], hours)),
            soc,
            label="储电量",
            color="#1f77b4",
            linewidth=2,
        )
        ax2.axhline(1200, color="#999999", linestyle=":")
        ax2.axhline(10800, color="#999999", linestyle=":")
        ax2.set_ylabel("储电量 (kWh)")
        ax2.set_ylim(0, 12000)
        ax.set_title(target.strftime("%Y-%m-%d"))
        ax.set_xlim(0, 24)
        ax.set_xticks(np.arange(0, 25, 3))
        lines1, labels1 = ax.get_legend_handles_labels()
        lines2, labels2 = ax2.get_legend_handles_labels()
        ax.legend(lines1 + lines2, labels1 + labels2, loc="upper left", fontsize=8)
    fig.suptitle("问题3指定日期：储能充放电与储电量", fontsize=15)
    fig.tight_layout(rect=(0, 0, 1, 0.97))
    save_figure(fig, output_path)
    plt.close(fig)


def plot_scenarios(scenarios: pd.DataFrame, output_path: Path) -> None:
    """绘制增加预报更新时点后的费用和紧急购电变化。"""
    aggregate = aggregate_forecast_scenarios(scenarios.to_dict(orient="records"))
    fig, ax1 = plt.subplots(figsize=(10, 5.5))
    x = np.arange(len(aggregate))
    axis_text_color = "#222222"
    ax1.bar(x - 0.18, aggregate["总费用_元"], width=0.36, label="总费用", color="#f5a684")
    ax1.set_ylabel("总费用 (元)", color=axis_text_color)
    ax1.tick_params(axis="y", labelcolor=axis_text_color)
    ax1.tick_params(axis="x", labelcolor=axis_text_color)
    ax1.set_xticks(x)
    ax1.set_xticklabels(aggregate["情景"])
    ax1.grid(axis="y", alpha=0.25)
    ax2 = ax1.twinx()
    ax2.bar(
        x + 0.18,
        aggregate["紧急购电量_kWh"],
        width=0.36,
        label="紧急购电量",
        color="#fcabed",
    )
    ax2.set_ylabel("紧急购电量 (kWh)", color=axis_text_color)
    ax2.tick_params(axis="y", labelcolor=axis_text_color)
    lines1, labels1 = ax1.get_legend_handles_labels()
    lines2, labels2 = ax2.get_legend_handles_labels()
    ax1.legend(lines1 + lines2, labels1 + labels2, loc="upper right")
    ax1.set_title("问题3：预报更新时点的边际价值")
    fig.tight_layout()
    save_figure(fig, output_path)
    plt.close(fig)


def write_report(
    output_path: Path,
    daily: pd.DataFrame,
    specified: pd.DataFrame,
    scenarios: pd.DataFrame,
    sensitivity: pd.DataFrame,
    storage,
    settlement_mode: str,
    scenario_count: int,
    scenario_lookback_days: int,
    window_days: int,
    terminal_soc_value_yuan_per_kwh: float,
) -> None:
    """写出问题3结果说明。"""
    period = daily[
        (daily["日期"].dt.date >= OUTPUT_START)
        & (daily["日期"].dt.date <= OUTPUT_END)
    ]
    aggregate = aggregate_forecast_scenarios(scenarios.to_dict(orient="records"))
    scenario_lookup = aggregate.set_index("情景")

    def improvement(before: str, after: str, column: str) -> tuple[float, float]:
        """返回增加更新时点后的绝对下降量和相对下降率。"""
        before_value = float(scenario_lookup.loc[before, column])
        after_value = float(scenario_lookup.loc[after, column])
        absolute = before_value - after_value
        relative = absolute / before_value * 100.0 if before_value > 0.0 else 0.0
        return absolute, relative

    cost_0_6, cost_0_6_pct = improvement("仅0:00预报", "更新至6:00", "总费用_元")
    cost_6_12, cost_6_12_pct = improvement("更新至6:00", "更新至12:00", "总费用_元")
    cost_12_18, cost_12_18_pct = improvement(
        "更新至12:00",
        "更新至18:00",
        "总费用_元",
    )
    emergency_0_6, emergency_0_6_pct = improvement(
        "仅0:00预报",
        "更新至6:00",
        "紧急购电量_kWh",
    )
    emergency_6_12, emergency_6_12_pct = improvement(
        "更新至6:00",
        "更新至12:00",
        "紧急购电量_kWh",
    )
    emergency_12_18, emergency_12_18_pct = improvement(
        "更新至12:00",
        "更新至18:00",
        "紧急购电量_kWh",
    )
    def update_verdict(absolute: float, relative: float) -> str:
        """根据费用改善方向生成更新时点结论。"""
        if absolute > 0.0 and relative >= 0.1:
            return "改善较明显，建议保留"
        if absolute > 0.0:
            return "改善很小，可按通信和计算成本决定是否保留"
        if abs(absolute) <= 1e-6:
            return "基本没有增量价值"
        return "费用略升，不建议只为此增加该更新时点"

    verdict_6 = update_verdict(cost_0_6, cost_0_6_pct)
    verdict_12 = update_verdict(cost_6_12, cost_6_12_pct)
    verdict_18 = update_verdict(cost_12_18, cost_12_18_pct)

    def csv_block(frame: pd.DataFrame) -> str:
        """用CSV文本展示表格，避免依赖可选的tabulate包。"""
        return frame.to_csv(index=False, float_format="%.6f").strip()

    lines = [
        "# 问题3结果说明",
        "",
        "## 模型口径",
        "",
        "- 附件3的“预报k小时”解释为发布时刻后第k小时的平均光伏功率，单位kW。",
        "- 每小时预报在6个10分钟区间内保持不变，电量按 功率×0.1666666667 h 计算。",
        "- 负荷基准仅使用当前日期之前的同星期历史实际曲线；历史不足时回退到"
        "此前最多7天均值，不读取当天未来实际负荷。",
        f"- 每天从此前{scenario_lookback_days}天预测残差中选取"
        f"{scenario_count}个代表情景，负荷和光伏误差成对进入优化。",
        "- 光伏情景以附件3对应发布时刻的实际预报为基准，并叠加同发布时刻、"
        "同提前期的历史预报误差。",
        f"- 情景生成参考{window_days}天跨日窗口，实时执行使用单位储能价值"
        f" {terminal_soc_value_yuan_per_kwh:.6f} 元/kWh。",
        "- 每天0:00使用0:00预报确定计划购电量g，g在当天剩余时段锁定。",
        "- 6:00、12:00、18:00只更新此后尚未执行时段的购电量q，不回改g。",
        "- 充放电量c、d按实际负荷和实际光伏逐10分钟实时调整，"
        "每个时刻只使用当前及过去真实数据。",
        "- 计划购电量高于调整购电量部分按交易时刻电价50%计违约费用。",
        "- 调整购电量高于计划购电量部分按交易时刻电价150%计费。",
        "- 最终实际光伏与预报的偏差由紧急购电或弃光结算。",
        "- 储能从2025-01-01 0:00的6000 kWh开始，跨日连续运行；"
        "前一日24:00储电量作为次日0:00初值，不要求每日回到6000 kWh。",
        f"- 费用结算口径：{settlement_mode}。",
        "",
        "## 储能参数",
        "",
        f"- 容量：{storage.capacity_kwh:.6f} kWh",
        f"- 最大充放电功率：{storage.power_kw:.6f} kW",
        f"- SOC范围：{storage.soc_min_kwh:.6f}~{storage.soc_max_kwh:.6f} kWh",
        f"- 充放电效率：{storage.efficiency:.6f}",
        "",
        "## 输出期汇总",
        "",
        "| 指标 | 数值 | 单位 |",
        "|---|---:|---|",
        f"| 计划购电量 | {period['计划购电量_kWh'].sum():.6f} | kWh |",
        f"| 调整购电量 | {period['调整购电量_kWh'].sum():.6f} | kWh |",
        f"| 紧急购电量 | {period['紧急购电量_kWh'].sum():.6f} | kWh |",
        f"| 计划购电费 | {period['计划购电费_元'].sum():.6f} | 元 |",
        f"| 调整费用 | {period['调整费用_元'].sum():.6f} | 元 |",
        f"| 紧急购电费 | {period['紧急购电费_元'].sum():.6f} | 元 |",
        f"| 总费用 | {period['总费用_元'].sum():.6f} | 元 |",
        "",
        "## 指定日期结果",
        "",
        "```text",
        csv_block(specified),
        "```",
        "",
        "## 是否需要增加预报更新时点",
        "",
        "```text",
        csv_block(aggregate),
        "```",
        "",
        "边际价值计算结果：",
        "",
        f"- 增加6:00预报：总费用下降 {cost_0_6:.6f} 元"
        f"（{cost_0_6_pct:.4f}%），紧急购电量下降 "
        f"{emergency_0_6:.6f} kWh（{emergency_0_6_pct:.4f}%）。",
        f"- 增加12:00预报：总费用再下降 {cost_6_12:.6f} 元"
        f"（{cost_6_12_pct:.4f}%），紧急购电量再下降 "
        f"{emergency_6_12:.6f} kWh（{emergency_6_12_pct:.4f}%）。",
        f"- 增加18:00预报：总费用再下降 {cost_12_18:.6f} 元"
        f"（{cost_12_18_pct:.6f}%），紧急购电量再下降 "
        f"{emergency_12_18:.6f} kWh（{emergency_12_18_pct:.6f}%）。",
        "",
        f"结论：6:00预报相对0:00预报{verdict_6}；"
        f"12:00相对6:00预报{verdict_12}；"
        f"18:00相对12:00预报{verdict_18}。",
        "",
        "仍需关注小时预报与10分钟实际光伏的误差，因为它会造成小时内功率不匹配。",
        "",
        "## 预报整体缩放灵敏度",
        "",
        "```text",
        csv_block(sensitivity),
        "```",
    ]
    output_path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    """问题3主流程。"""
    configure_console()
    args = parse_args()
    if args.scenarios <= 0:
        raise ValueError("情景数量必须为正整数。")
    if args.scenario_lookback_days <= 0:
        raise ValueError("历史误差回看天数必须为正整数。")
    if args.window_days <= 0:
        raise ValueError("滚动窗口天数必须为正整数。")
    if args.terminal_value_factor < 0.0:
        raise ValueError("终端储能价值倍率不能为负。")
    if args.scenario_time_limit <= 0.0:
        raise ValueError("情景MILP时间上限必须为正。")
    script_dir = Path(__file__).resolve().parent
    p2 = load_problem2_module()
    inputs = locate_inputs(script_dir)
    storage = p2.read_storage_parameters(inputs["pdf"])
    forecasts = read_attachment3(inputs["attachment3"])

    base_price = p2.read_price_curve(inputs["attachment1"])
    fallback_load_profile = read_attachment1_load_energy(inputs["attachment1"])
    all_dates = pd.date_range("2025-01-01", "2025-12-31", freq="D").date
    price_by_date = {current_date: base_price.copy() for current_date in all_dates}
    data = prepare_actual_data(p2, inputs["attachment2"], price_by_date)
    print("附件路径：")
    for key, value in inputs.items():
        print(f"{key} = {value}")
    print("储能参数：")
    print(
        f"容量={storage.capacity_kwh:.6f} kWh，"
        f"最大功率={storage.power_kw:.6f} kW，"
        f"初始SOC={storage.initial_kwh:.6f} kWh，"
        f"效率={storage.efficiency:.6f}。"
    )
    print_quantity_checks(data, forecasts, "附件1电价")
    terminal_soc_value = (
        args.terminal_value_factor
        * float(np.mean(base_price[:30]) / storage.efficiency)
    )

    # 所有代码和输出均位于“问题三”目录内；脚本可从任意当前目录启动。
    output_dir = (
        args.output_dir.expanduser().resolve()
        if args.output_dir is not None
        else script_dir / "results"
    )
    tables_dir = output_dir / "tables"
    figures_dir = output_dir / "figures"
    logs_dir = output_dir / "logs"
    for directory in (output_dir, tables_dir, figures_dir, logs_dir):
        directory.mkdir(parents=True, exist_ok=True)

    baseline = run_baseline_benchmark(p2, data, storage)
    detail, daily, scenarios = solve_problem3(
        p2,
        data,
        forecasts,
        storage,
        args.settlement_mode,
        fallback_load_profile_kwh=fallback_load_profile,
        scenario_count=args.scenarios,
        scenario_lookback_days=args.scenario_lookback_days,
        window_days=args.window_days,
        terminal_soc_value_yuan_per_kwh=terminal_soc_value,
        scenario_time_limit_s=args.scenario_time_limit,
    )
    specified = summarize_specified_dates(daily, include_adjustment=True)
    validation = validate_result_detail(detail, storage, include_adjustment=True)
    print("问题3约束校验：")
    for key, value in validation.items():
        print(f"{key} = {value:.10f}")

    sensitivity = run_forecast_sensitivity(
        p2,
        data,
        forecasts,
        storage,
        args.settlement_mode,
        initial_soc_by_date={
            row["日期"].date(): float(row["0:00储电量_kWh"])
            for _, row in daily.iterrows()
        },
        scenario_count=args.scenarios,
        scenario_lookback_days=args.scenario_lookback_days,
        window_days=args.window_days,
        terminal_soc_value_yuan_per_kwh=terminal_soc_value,
        scenario_time_limit_s=args.scenario_time_limit,
        fallback_profile_kwh=fallback_load_profile,
    )
    scenario_summary = aggregate_forecast_scenarios(
        scenarios.to_dict(orient="records")
    )
    paper_tables = write_paper_tables(
        p2,
        detail,
        storage,
        output_dir,
    )

    result_path = output_dir / "result3.xlsx"
    write_official_result(
        p2,
        inputs["result3"],
        result_path,
        detail,
        storage,
        include_adjustment=True,
    )
    write_specified_date_workbook(
        specified,
        output_dir / "指定日期结果.xlsx",
        include_adjustment=True,
    )
    detail.to_csv(
        tables_dir / "逐10分钟计划调整明细.csv",
        index=False,
        encoding="utf-8-sig",
    )
    daily.to_csv(tables_dir / "逐日汇总.csv", index=False, encoding="utf-8-sig")
    specified.to_csv(
        tables_dir / "指定日期结果.csv",
        index=False,
        encoding="utf-8-sig",
    )
    scenarios.to_csv(
        tables_dir / "预报更新情景逐日.csv",
        index=False,
        encoding="utf-8-sig",
    )
    scenario_summary.to_csv(
        tables_dir / "预报更新情景汇总.csv",
        index=False,
        encoding="utf-8-sig",
    )
    sensitivity.to_csv(
        tables_dir / "预报灵敏度分析.csv",
        index=False,
        encoding="utf-8-sig",
    )
    baseline.to_csv(
        tables_dir / "基准算例_储能不动作.csv",
        index=False,
        encoding="utf-8-sig",
    )
    paper_tables["表3明细"].to_csv(
        tables_dir / "表3_指定日期紧急购电量.csv",
        index=False,
        encoding="utf-8-sig",
    )
    plot_forecast_and_dispatch(
        detail,
        figures_dir / "指定日期_预报更新与购电调整.png",
    )
    plot_storage(
        detail,
        storage,
        figures_dir / "指定日期_储能充放电与储电量.png",
    )
    plot_scenarios(
        scenarios,
        figures_dir / "预报更新时点_边际价值.png",
    )
    write_report(
        output_dir / "问题三_结果说明.md",
        daily,
        specified,
        scenarios,
        sensitivity,
        storage,
        args.settlement_mode,
        args.scenarios,
        args.scenario_lookback_days,
        args.window_days,
        terminal_soc_value,
    )
    summary = {
        "输出期": {
            "开始": OUTPUT_START.isoformat(),
            "结束": OUTPUT_END.isoformat(),
            "天数": int(len(daily)),
            "计划购电量_kWh": float(daily["计划购电量_kWh"].sum()),
            "调整购电量_kWh": float(daily["调整购电量_kWh"].sum()),
            "紧急购电量_kWh": float(daily["紧急购电量_kWh"].sum()),
            "总费用_元": float(daily["总费用_元"].sum()),
        },
        "费用结算口径": args.settlement_mode,
        "储能边界口径": "2025-01-01至12-31跨日连续，前一日最终SOC作为次日初值",
        "储能执行口径": "0:00锁定计划购电g，预报点更新调整购电量q，充放电量按实际数据实时调整",
        "购电修改规则": "g仅允许在0:00修改；q仅允许在6:00、12:00、18:00修改",
        "负荷预测口径": "当前日期之前同星期历史实际曲线，并叠加历史预测误差情景",
        "光伏预测口径": "附件3对应发布时刻预报，并叠加历史同提前期预报误差情景",
        "情景数量": args.scenarios,
        "历史误差回看天数": args.scenario_lookback_days,
        "滚动窗口天数": args.window_days,
        "终端储能价值_元每kWh": terminal_soc_value,
        "计划阶段使用当天未来实际负荷": False,
        "约束校验": validation,
        "指定日期结果": specified.assign(
            日期=specified["日期"].dt.strftime("%Y-%m-%d")
        ).to_dict(orient="records"),
    }
    (tables_dir / "summary.json").write_text(
        json.dumps(summary, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(f"result3.xlsx = {result_path}")
    for name in ("表1", "表2", "表3"):
        print(f"{name} = {paper_tables[name]}")
    print(f"结果目录 = {output_dir}")
    print("问题3处理完成。")


if __name__ == "__main__":
    main()

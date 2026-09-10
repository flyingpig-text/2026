from __future__ import annotations

from pathlib import Path

import openpyxl
from pypdf import PdfReader


SCRIPT_DIR = Path(__file__).resolve().parent
ATTACHMENT_DIR = Path(r"D:\46884\Documents\2026\题目\附件")


def show_workbook(path: Path, preview_rows: int = 8, preview_columns: int = 18) -> None:
    """打印工作簿的工作表、维度和前几行，不修改附件。"""
    print("=" * 100)
    print(f"文件: {path}")
    print(f"存在: {path.exists()} | 大小: {path.stat().st_size if path.exists() else 0} B")
    if not path.exists():
        return

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    print(f"工作表: {workbook.sheetnames}")
    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        print("-" * 100)
        print(
            f"工作表: {sheet_name!r} | "
            f"行数: {sheet.max_row} | 列数: {sheet.max_column}"
        )
        for row_number, row in enumerate(
            sheet.iter_rows(
                min_row=1,
                max_row=min(preview_rows, sheet.max_row),
                max_col=min(preview_columns, sheet.max_column),
                values_only=True,
            ),
            start=1,
        ):
            values = ["" if value is None else repr(value) for value in row]
            suffix = " | ..." if sheet.max_column > preview_columns else ""
            print(f"{row_number:>4}: " + " | ".join(values) + suffix)
    workbook.close()


def show_pdf_sections(path: Path) -> None:
    """打印 PDF 中与问题 2、附录 1、表 3 有关的页面。"""
    print("=" * 100)
    print(f"文件: {path}")
    reader = PdfReader(path)
    keywords = ("问题2", "问题 2", "附录1", "附录 1", "表3", "表 3", "储能")
    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if any(keyword in text for keyword in keywords):
            print("-" * 100)
            print(f"PDF 第 {page_number} 页")
            print(text)


def main() -> None:
    """检查问题 2 涉及的工作簿和结果模板。"""
    paths = [
        ATTACHMENT_DIR / "附件1.xlsx",
        ATTACHMENT_DIR / "附件2.xlsx",
        ATTACHMENT_DIR / "附件4.xlsx",
        ATTACHMENT_DIR / "附件5" / "result2.xlsx",
        ATTACHMENT_DIR / "附件5" / "result4-2.xlsx",
    ]
    print(f"脚本目录: {SCRIPT_DIR}")
    print(f"附件目录: {ATTACHMENT_DIR}")
    for path in paths:
        show_workbook(path)
    show_pdf_sections(ATTACHMENT_DIR.parent / "C题.pdf")


if __name__ == "__main__":
    main()
